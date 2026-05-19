import fitz  # PyMuPDF
import pytesseract
from PIL import Image
import io
import json

pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

IMAGE_EXTENSIONS = {'.jpg', '.jpeg', '.png', '.bmp', '.tiff', '.webp'}
CODE_EXTENSIONS = {'.py', '.php', '.js', '.ts', '.html', '.css', '.java', '.c', '.cpp'}

class AIService:
    @staticmethod
    def extract_text(file_content: bytes, filename: str) -> str:
        text = ""
        ext = '.' + filename.lower().rsplit('.', 1)[-1]
        try:
            # PDF
            if ext == '.pdf':
                doc = fitz.open(stream=file_content, filetype="pdf")
                for page in doc:
                    text += page.get_text()
                doc.close()

            # Images → OCR
            elif ext in IMAGE_EXTENSIONS:
                image = Image.open(io.BytesIO(file_content))
                text = pytesseract.image_to_string(image, lang='fra')

            # Word (.docx)
            elif ext == '.docx':
                import docx
                doc = docx.Document(io.BytesIO(file_content))
                text = "\n".join([p.text for p in doc.paragraphs])

            # PowerPoint (.pptx)
            elif ext == '.pptx':
                from pptx import Presentation
                prs = Presentation(io.BytesIO(file_content))
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, 'text'):
                            text += shape.text + "\n"

            # Texte brut
            elif ext == '.txt':
                text = file_content.decode('utf-8', errors='ignore')

            # CSV
            elif ext == '.csv':
                text = file_content.decode('utf-8', errors='ignore')

            # JSON + Jupyter (.ipynb)
            elif ext in {'.json', '.ipynb'}:
                data = json.loads(file_content.decode('utf-8', errors='ignore'))
                text = json.dumps(data, ensure_ascii=False, indent=2)

            # Code source
            elif ext in CODE_EXTENSIONS:
                text = file_content.decode('utf-8', errors='ignore')

            else:
                print(f"[AIService] Format non supporté : {ext}")
                return ""

            return text.strip()

        except Exception as e:
            print(f"Erreur d'extraction : {e}")
            return ""